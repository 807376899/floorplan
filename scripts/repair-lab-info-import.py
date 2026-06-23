from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "lab-info-import"


def text(value) -> str:
  return "" if value is None else str(value).strip()


def digits(value) -> str:
  return re.sub(r"\D+", "", text(value))


def load_rows(sheet):
  headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
  rows = []
  for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
    rows.append((index, {headers[i]: values[i] for i in range(len(headers))}))
  return headers, rows


def find_source_workbook() -> Path:
  matches = [
    path for path in ROOT.glob("*.xlsx")
    if path.name != "floor-room-baseline-upload-package.xlsx" and not path.name.startswith("~$")
  ]
  if not matches:
    raise FileNotFoundError("未找到源文件：实验室信息.xlsx")
  return matches[0]


def find_template_workbook() -> Path:
  matches = [
    path for path in OUTPUT_DIR.glob("*.xlsx")
    if not path.name.startswith("~$") and "-修复" not in path.stem
  ]
  if not matches:
    matches = [path for path in OUTPUT_DIR.glob("*.xlsx") if not path.name.startswith("~$")]
  if not matches:
    raise FileNotFoundError("未找到 outputs/lab-info-import 下的导入模板")
  return matches[0]


def sheet2_room_by_name(source_path: Path) -> dict[str, str]:
  workbook = load_workbook(source_path, read_only=True, data_only=True)
  sheet = workbook["Sheet2"]
  rows = list(sheet.iter_rows(values_only=True))
  headers = [text(value) for value in rows[0]]
  name_index = headers.index("*实训室名称")
  room_index = headers.index("*房间号")
  result = {}
  for row in rows[1:]:
    name = text(row[name_index])
    room = text(row[room_index])
    if name and room and name not in result:
      result[name] = room
  return result


def sheet1_floor_by_name_doors(source_path: Path) -> dict[tuple[str, str, str], str]:
  workbook = load_workbook(source_path, read_only=True, data_only=True)
  sheet = workbook["sheet1"]
  headers = [text(cell.value) for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
  name_index = headers.index("场地名称")
  floor_index = headers.index("楼层")
  front_index = headers.index("前门牌")
  rear_index = headers.index("后门牌")
  result = {}
  for row in sheet.iter_rows(min_row=5, values_only=True):
    name = text(row[name_index])
    front = text(row[front_index])
    rear = text(row[rear_index])
    floor = text(row[floor_index])
    if name and front:
      result[(name, front, rear)] = floor
  return result


def unique_code(base: str, used: set[str]) -> str:
  code = base
  suffix = 2
  while code in used:
    code = f"{base}-{suffix:03d}"
    suffix += 1
  used.add(code)
  return code


def repaired_space_code(row: dict) -> str:
  building_digits = digits(row.get("building_code"))[-2:]
  floor_digits = digits(row.get("floor_code")) or "0"
  front_digits = digits(row.get("front_door")) or digits(row.get("space_code"))
  rear_digits = digits(row.get("rear_door")) or front_digits
  return f"001{building_digits}{floor_digits.zfill(2)}{front_digits}{rear_digits}"


def repair_workbook(source_path: Path, template_path: Path) -> dict[str, int | str]:
  sheet2_rooms = sheet2_room_by_name(source_path)
  sheet1_floors = sheet1_floor_by_name_doors(source_path)
  workbook = load_workbook(template_path)
  spaces_sheet = workbook["spaces"]
  labs_sheet = workbook["labs"]
  assignments_sheet = workbook["plan_assignments"]

  space_headers, space_rows = load_rows(spaces_sheet)
  lab_headers, lab_rows = load_rows(labs_sheet)
  assignment_headers, assignment_rows = load_rows(assignments_sheet)
  space_col = {header: index + 1 for index, header in enumerate(space_headers)}
  lab_col = {header: index + 1 for index, header in enumerate(lab_headers)}
  assignment_col = {header: index + 1 for index, header in enumerate(assignment_headers)}

  labs_by_name = defaultdict(list)
  for row_index, row in lab_rows:
    labs_by_name[text(row.get("lab_name"))].append((row_index, row))

  assignments_by_lab = defaultdict(list)
  for row_index, row in assignment_rows:
    assignments_by_lab[text(row.get("lab_code"))].append((row_index, row))

  groups = defaultdict(list)
  for row_index, row in space_rows:
    groups[text(row.get("space_code"))].append((row_index, row))

  used_space_codes = {text(row.get("space_code")) for _, row in space_rows if text(row.get("space_code"))}
  used_lab_codes = {text(row.get("lab_code")) for _, row in lab_rows if text(row.get("lab_code"))}
  updates = 0
  for code, rows in groups.items():
    physical_keys = {
      (
        text(row.get("building_code")),
        text(row.get("floor_code")),
        text(row.get("front_door")),
        text(row.get("rear_door")),
      )
      for _, row in rows
    }
    if len(rows) < 2 or len(physical_keys) < 2:
      continue

    for row_index, row in rows:
      room_name = text(row.get("space_name"))
      source_floor = sheet1_floors.get((room_name, text(row.get("front_door")), text(row.get("rear_door"))))
      if source_floor and source_floor != text(row.get("floor_code")):
        spaces_sheet.cell(row_index, space_col["floor_code"]).value = source_floor
        row["floor_code"] = source_floor
      if sheet2_rooms.get(room_name) == code:
        for lab_index, lab in labs_by_name.get(room_name, []):
          old_lab_code = text(lab.get("lab_code"))
          target_lab_code = code.removeprefix("001")
          if old_lab_code != target_lab_code:
            used_lab_codes.discard(old_lab_code)
            used_lab_codes.add(target_lab_code)
            labs_sheet.cell(lab_index, lab_col["lab_code"]).value = target_lab_code
            for assignment_index, _assignment in assignments_by_lab.get(old_lab_code, []):
              assignments_sheet.cell(assignment_index, assignment_col["lab_code"]).value = target_lab_code
            updates += 1
        continue

      old_space_code = code
      used_space_codes.discard(old_space_code)
      new_space_code = unique_code(repaired_space_code(row), used_space_codes)
      spaces_sheet.cell(row_index, space_col["space_code"]).value = new_space_code

      for lab_index, lab in labs_by_name.get(room_name, []):
        old_lab_code = text(lab.get("lab_code"))
        new_lab_code = unique_code(new_space_code.removeprefix("001"), used_lab_codes)
        labs_sheet.cell(lab_index, lab_col["lab_code"]).value = new_lab_code
        for assignment_index, assignment in assignments_by_lab.get(old_lab_code, []):
          if text(assignment.get("space_code")) == old_space_code:
            assignments_sheet.cell(assignment_index, assignment_col["space_code"]).value = new_space_code
          assignments_sheet.cell(assignment_index, assignment_col["lab_code"]).value = new_lab_code
        updates += 1

  saved_path = template_path
  try:
    workbook.save(saved_path)
  except PermissionError:
    saved_path = template_path.with_name(f"{template_path.stem}-修复{template_path.suffix}")
    workbook.save(saved_path)
  return {"updates": updates, "saved_path": str(saved_path)}


def main() -> None:
  source_path = find_source_workbook()
  template_path = find_template_workbook()
  result = repair_workbook(source_path, template_path)
  print(f"repaired {result['saved_path']} using {source_path}; updates={result['updates']}")


if __name__ == "__main__":
  main()
